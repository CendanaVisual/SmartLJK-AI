import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_report(exam_data, results, questions, stats, question_analysis):
    """
    Generate Excel report for exam results.
    """
    wb = Workbook()
    
    # Border style
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # --- Sheet 1: Rekap Nilai Siswa ---
    ws1 = wb.active
    ws1.title = "Rekap Nilai Siswa"
    
    headers_ws1 = ["No", "NIS", "Nama Siswa", "Kelas", "Total Skor", "Nilai Akhir", "Status", "Akurasi OMR"]
    for col, header in enumerate(headers_ws1, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    for i, res in enumerate(results, 1):
        row_idx = i + 1
        
        status = res.get("status", "Remedial")
        status_color = "00B050" if status.lower() == "lulus" else "FF0000"
        
        row_data = [
            i,
            res.get("student_id_number", ""),
            res.get("student_name", ""),
            res.get("class_name", ""),
            res.get("total_score", 0),
            res.get("percentage", 0),
            status,
            res.get("omr_confidence", 0)
        ]
        
        fill_color = "FFFFFF" if i % 2 != 0 else "D9E2F3"
        fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        
        for col, val in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col, value=val)
            cell.fill = fill
            cell.border = thin_border
            if col == 7: # Status
                cell.font = Font(color=status_color)
            if col in [1, 5, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="center")
                
    # Auto-fit columns
    for col in ws1.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws1.column_dimensions[column].width = adjusted_width

    # --- Sheet 2: Statistik Kelas ---
    ws2 = wb.create_sheet(title="Statistik Kelas")
    
    ws2["A1"] = "STATISTIK KELAS"
    ws2["A1"].font = Font(bold=True, size=14)
    
    stat_rows = [
        (3, "Nama Ujian:", exam_data.get("title", "")),
        (4, "Mata Pelajaran:", exam_data.get("subject", "")),
        (5, "Kelas:", exam_data.get("class_name", "")),
        (6, "KKM:", exam_data.get("passing_score", "")),
        (8, "Jumlah Peserta:", stats.get("total_students", 0)),
        (9, "Rata-rata Kelas:", f"{stats.get('avg_score', 0):.2f}"),
        (10, "Nilai Tertinggi:", stats.get("max_score", 0)),
        (11, "Nilai Terendah:", stats.get("min_score", 0)),
        (12, "Jumlah Lulus:", stats.get("pass_count", 0)),
        (13, "Jumlah Remedial:", stats.get("fail_count", 0)),
    ]
    
    total = stats.get("total_students", 0)
    pass_pct = (stats.get("pass_count", 0) / total * 100) if total > 0 else 0
    stat_rows.append((14, "Tingkat Kelulusan:", f"{pass_pct:.1f}%"))
    
    for r, label, val in stat_rows:
        cell_label = ws2.cell(row=r, column=1, value=label)
        cell_label.font = Font(bold=True)
        cell_val = ws2.cell(row=r, column=2, value=val)
        cell_val.alignment = Alignment(horizontal="right")
        
    ws2.column_dimensions['A'].width = 20
    ws2.column_dimensions['B'].width = 30

    # --- Sheet 3: Analisis Butir Soal ---
    ws3 = wb.create_sheet(title="Analisis Butir Soal")
    
    headers_ws3 = ["No. Soal", "Tipe Soal", "Jumlah Benar", "Jumlah Salah", "Total Peserta", "Persentase Benar", "Kategori"]
    for col, header in enumerate(headers_ws3, 1):
        cell = ws3.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    for i, qa in enumerate(question_analysis, 1):
        row_idx = i + 1
        pct = qa.get("percentage", 0)
        
        if pct > 75:
            kategori = "Mudah"
            cat_color = "92D050" # Light Green
        elif pct >= 25:
            kategori = "Sedang"
            cat_color = "FFC000" # Yellow
        else:
            kategori = "Sulit"
            cat_color = "FF0000" # Red
            
        total_q = qa.get("total_count", 0)
        correct_q = qa.get("correct_count", 0)
        wrong_q = total_q - correct_q
        
        row_data = [
            qa.get("question_number", i),
            qa.get("question_type", ""),
            correct_q,
            wrong_q,
            total_q,
            pct / 100.0, # format as percentage in excel
            kategori
        ]
        
        for col, val in enumerate(row_data, 1):
            cell = ws3.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
            if col == 6:
                cell.number_format = '0.00%'
            if col == 7:
                cell.fill = PatternFill(start_color=cat_color, end_color=cat_color, fill_type="solid")
                if cat_color == "FF0000":
                    cell.font = Font(color="FFFFFF")
                    
    for col in ws3.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws3.column_dimensions[column].width = max_length + 2

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
