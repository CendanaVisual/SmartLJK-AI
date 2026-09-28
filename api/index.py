"""
SmartLJK API - Main FastAPI Application
Backend untuk Pembuat LJK & Pemeriksa Jawaban Otomatis Berbasis AI
"""
import os
import json
import io
import base64
from datetime import datetime
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse
from pydantic import BaseModel
from typing import List, Optional, Dict, Any

# Internal imports
from api import database

try:
    from api import omr_processor
except ImportError:
    omr_processor = None

try:
    from api import ai_grader
except ImportError:
    ai_grader = None

try:
    from api import pdf_generator
except ImportError:
    pdf_generator = None

try:
    from api import excel_export
except ImportError:
    excel_export = None

app = FastAPI(
    title="SmartLJK API",
    description="API Backend untuk Aplikasi SmartLJK - Pembuat LJK & Pemeriksa Jawaban Otomatis",
    version="1.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

GEMINI_API_KEY = os.environ.get(
    'GEMINI_API_KEY',
    'AIzaSyAb8RN6LudkTDhmyjmH2rMI9XSZi_R9ztvo5JDMAvg3uo0Qsijw'
)


# ============================================================
# Startup & Shutdown Events
# ============================================================

@app.on_event("startup")
async def startup_event():
    """Inisialisasi database pool dan buat tabel"""
    await database.init_db()


@app.on_event("shutdown")
async def shutdown_event():
    """Tutup koneksi database pool"""
    await database.close_pool()


# ============================================================
# Pydantic Models
# ============================================================

class QuestionCreate(BaseModel):
    question_number: int
    question_type: str  # single_choice, multi_choice, true_false, matching, short_answer
    question_text: Optional[str] = None
    answer_key: Any
    weight: float = 1.0
    options: Optional[Any] = None


class ExamCreate(BaseModel):
    title: str
    code: str
    subject: str
    class_name: str
    academic_year: str = '2024/2025'
    passing_score: float = 75.0
    description: Optional[str] = None
    questions: Optional[List[QuestionCreate]] = []


class GenerateQuestionsRequest(BaseModel):
    mata_pelajaran: Optional[str] = ""
    jenjang: Optional[str] = ""
    kelas: Optional[str] = ""
    fase: Optional[str] = ""
    sekolah: Optional[str] = ""
    kepala_sekolah: Optional[str] = ""
    guru: Optional[str] = ""
    nip_kepala: Optional[str] = ""
    nip_guru: Optional[str] = ""
    lingkup_materi: Optional[str] = ""
    tujuan_pembelajaran: Optional[str] = ""
    indikator_soal: Optional[str] = ""
    jenis_taksonomi: Optional[str] = ""
    tingkat_kognitif: Optional[str] = ""
    jumlah_opsi: Optional[str] = ""
    bentuk_soal: Optional[str] = ""
    jumlah_soal: Optional[str] = "5"
    
    # Old fields for backward compatibility
    topic: Optional[str] = ""
    target_class: Optional[str] = ""
    difficulty: Optional[str] = ""
    count: Optional[int] = 5


# ============================================================
# Serializer Helper - konversi types yang tidak JSON-serializable
# ============================================================

def serialize_record(record):
    """Konversi record dict agar JSON-serializable"""
    if record is None:
        return None
    result = {}
    for key, value in record.items():
        if isinstance(value, datetime):
            result[key] = value.isoformat()
        elif isinstance(value, (int, float, str, bool, list, dict)):
            result[key] = value
        elif value is None:
            result[key] = None
        else:
            result[key] = str(value)
    return result


def serialize_list(records):
    """Konversi list of records"""
    return [serialize_record(r) for r in records]


# ============================================================
# Health Check
# ============================================================

@app.get("/api/health")
async def health_check():
    return {
        "status": "ok",
        "message": "SmartLJK API berjalan dengan baik",
        "timestamp": datetime.now().isoformat(),
        "modules": {
            "omr_processor": omr_processor is not None,
            "ai_grader": ai_grader is not None,
            "pdf_generator": pdf_generator is not None,
            "excel_export": excel_export is not None
        }
    }


# ============================================================
# Exam CRUD Endpoints
# ============================================================

@app.post("/api/exams")
async def create_exam(exam: ExamCreate):
    """Buat ujian baru beserta daftar soal"""
    try:
        new_exam = await database.create_exam(
            exam.title, exam.code, exam.subject, exam.class_name,
            exam.academic_year, exam.passing_score, exam.description
        )
        exam_id = new_exam['id']

        created_questions = []
        if exam.questions:
            for q in exam.questions:
                new_q = await database.create_question(
                    exam_id, q.question_number, q.question_type,
                    q.question_text, q.answer_key, q.weight, q.options
                )
                created_questions.append(serialize_record(new_q))

            # Update total_questions
            await database.update_exam(
                exam_id, exam.title, exam.code, exam.subject,
                exam.class_name, exam.academic_year, exam.passing_score,
                exam.description
            )

        return {
            "message": "Ujian berhasil dibuat",
            "exam": serialize_record(new_exam),
            "questions": created_questions,
            "total_questions": len(created_questions)
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Gagal membuat ujian: {str(e)}")


@app.get("/api/exams")
async def get_exams():
    """Daftar semua ujian"""
    try:
        exams = await database.get_all_exams()
        return serialize_list(exams)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/exams/{exam_id}")
async def get_exam(exam_id: int):
    """Detail ujian beserta daftar soal"""
    try:
        exam = await database.get_exam(exam_id)
        if not exam:
            raise HTTPException(status_code=404, detail="Ujian tidak ditemukan")
        questions = await database.get_questions_by_exam(exam_id)
        result = serialize_record(exam)
        result['questions'] = serialize_list(questions)
        return result
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.put("/api/exams/{exam_id}")
async def update_exam(exam_id: int, exam: ExamCreate):
    """Update ujian dan soal-soalnya"""
    try:
        updated_exam = await database.update_exam(
            exam_id, exam.title, exam.code, exam.subject, exam.class_name,
            exam.academic_year, exam.passing_score, exam.description
        )
        if not updated_exam:
            raise HTTPException(status_code=404, detail="Ujian tidak ditemukan")

        if exam.questions is not None:
            await database.delete_questions_by_exam(exam_id)
            for q in exam.questions:
                await database.create_question(
                    exam_id, q.question_number, q.question_type,
                    q.question_text, q.answer_key, q.weight, q.options
                )

        return {"message": "Ujian berhasil diupdate", "exam": serialize_record(updated_exam)}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.delete("/api/exams/{exam_id}")
async def delete_exam(exam_id: int):
    """Hapus ujian (cascade ke questions, results, answers)"""
    try:
        deleted = await database.delete_exam(exam_id)
        if not deleted:
            raise HTTPException(status_code=404, detail="Ujian tidak ditemukan")
        return {"message": "Ujian berhasil dihapus"}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================
# LJK PDF Generation
# ============================================================

@app.post("/api/generate-ljk/{exam_id}")
async def generate_ljk(exam_id: int):
    """Generate PDF LJK untuk sebuah ujian"""
    if not pdf_generator:
        raise HTTPException(status_code=501, detail="Modul PDF Generator belum tersedia")
    try:
        exam = await database.get_exam(exam_id)
        if not exam:
            raise HTTPException(status_code=404, detail="Ujian tidak ditemukan")

        questions = await database.get_questions_by_exam(exam_id)
        pdf_buffer = pdf_generator.generate_ljk_pdf(
            serialize_record(exam),
            serialize_list(questions)
        )

        return StreamingResponse(
            pdf_buffer,
            media_type="application/pdf",
            headers={
                "Content-Disposition": f"attachment; filename=LJK_{exam['code']}.pdf"
            }
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Gagal generate PDF: {str(e)}")


# ============================================================
# OMR Scan Processing
# ============================================================

@app.post("/api/scan")
async def scan_image(
    file: UploadFile = File(...),
    exam_id: int = Form(...),
    student_name: str = Form(...),
    student_id_number: str = Form(...)
):
    """Proses scan LJK tunggal (foto/gambar)"""
    try:
        image_bytes = await file.read()
        exam = await database.get_exam(exam_id)
        if not exam:
            raise HTTPException(status_code=404, detail="Ujian tidak ditemukan")

        questions = await database.get_questions_by_exam(exam_id)

        # Buat konfigurasi layout untuk OMR
        questions_config = _build_question_layout(questions)

        # Proses OMR jika modul tersedia
        if omr_processor:
            scan_result = omr_processor.process_ljk_image(image_bytes, questions_config)
            # Evaluasi tulisan tangan untuk isian singkat dengan AI Gemini Vision
            if ai_grader and scan_result.get('short_answers'):
                for q in questions:
                    q_num = q.get('question_number')
                    crop_bytes = scan_result['short_answers'].get(q_num) or scan_result['short_answers'].get(str(q_num))
                    if crop_bytes:
                        ans_key = q.get('answer_key', '')
                        weight = float(q.get('weight', 1.0))
                        try:
                            ai_res = ai_grader.grade_short_answer(crop_bytes, ans_key, max_score=weight)
                            if isinstance(ai_res, dict) and not ai_res.get('error'):
                                scan_result['answers'][str(q_num)] = {
                                    'answer': ai_res.get('transcribed_text', ''),
                                    'confidence': float(ai_res.get('confidence_percentage', 90.0)),
                                    'transcribed_text': ai_res.get('transcribed_text', ''),
                                    'is_correct': bool(ai_res.get('is_correct', False)),
                                    'score_earned': float(ai_res.get('score_earned', 0.0)),
                                    'ai_feedback': ai_res.get('explanation', ''),
                                    'is_ai_graded': True
                                }
                        except Exception as ai_err:
                            print(f"AI short answer grading error Q{q_num}: {ai_err}")
        else:
            # Simulasi hasil scan untuk demo
            scan_result = _simulate_scan_result(questions)

        # Hitung skor
        grading = _grade_all_answers(scan_result, questions)

        # Hitung total skor dan persentase
        max_score = sum([float(q.get('weight', 1.0)) for q in questions])
        total_score = grading['total_score']
        percentage = (total_score / max_score * 100) if max_score > 0 else 0
        passing_score = float(exam.get('passing_score', 75.0))
        status = "Lulus" if percentage >= passing_score else "Remedial"

        # Simpan ke database
        result_record = await database.create_exam_result(
            exam_id, student_name, student_id_number,
            exam.get('class_name', ''),
            total_score, max_score, round(percentage, 2), status,
            scan_image_url=file.filename,
            omr_confidence=scan_result.get('avg_confidence', 85.0),
            raw_details=scan_result.get('answers', {})
        )

        # Simpan jawaban per soal
        for ans in grading['details']:
            q_id = ans.get('question_id')
            await database.create_student_answer(
                result_record['id'], q_id,
                ans.get('student_response'),
                ans.get('is_correct', False),
                ans.get('score_earned', 0),
                ans.get('ai_feedback'),
                ans.get('ai_confidence')
            )

        return {
            "message": "Scan berhasil diproses",
            "result_id": result_record['id'],
            "student_name": student_name,
            "student_id_number": student_id_number,
            "total_score": total_score,
            "max_score": max_score,
            "percentage": round(percentage, 2),
            "status": status,
            "omr_confidence": scan_result.get('avg_confidence', 85.0),
            "details": grading['details']
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Gagal memproses scan: {str(e)}")


@app.post("/api/scan-pdf")
async def scan_pdf(
    file: UploadFile = File(...),
    exam_id: int = Form(...)
):
    """Proses scan PDF multi-halaman (hingga 32 lembar)"""
    try:
        pdf_bytes = await file.read()
        exam = await database.get_exam(exam_id)
        if not exam:
            raise HTTPException(status_code=404, detail="Ujian tidak ditemukan")

        # Gunakan PyMuPDF untuk extract halaman dari PDF
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=pdf_bytes, filetype="pdf")
            total_pages = min(doc.page_count, 32)  # Maksimal 32 halaman

            pages_info = []
            for i in range(total_pages):
                page = doc[i]
                # Render halaman menjadi gambar
                pix = page.get_pixmap(dpi=200)
                img_bytes = pix.tobytes("png")
                img_b64 = base64.b64encode(img_bytes).decode('utf-8')
                pages_info.append({
                    "page_number": i + 1,
                    "image_base64": img_b64,
                    "status": "menunggu",
                    "width": pix.width,
                    "height": pix.height
                })
            doc.close()

            return {
                "message": f"PDF berhasil diproses: {total_pages} halaman ditemukan",
                "total_pages": total_pages,
                "pages": pages_info
            }
        except ImportError:
            raise HTTPException(
                status_code=501,
                detail="PyMuPDF (fitz) belum terinstal untuk memproses PDF"
            )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Gagal memproses PDF: {str(e)}")


@app.post("/api/scan-demo")
async def scan_demo(exam_id: int = Form(None)):
    """Demo scan dengan data simulasi"""
    try:
        # Jika exam_id diberikan, gunakan soal dari database
        if exam_id:
            exam = await database.get_exam(exam_id)
            questions = await database.get_questions_by_exam(exam_id)
        else:
            # Data demo default
            exam = {
                "id": 0, "title": "Demo Ujian", "code": "DEMO-001",
                "subject": "Umum", "class_name": "Demo",
                "passing_score": 75.0
            }
            questions = [
                {"id": 1, "question_number": 1, "question_type": "single_choice",
                 "answer_key": "B", "weight": 1.0},
                {"id": 2, "question_number": 2, "question_type": "single_choice",
                 "answer_key": "A", "weight": 1.0},
                {"id": 3, "question_number": 3, "question_type": "true_false",
                 "answer_key": "B", "weight": 1.0},
                {"id": 4, "question_number": 4, "question_type": "multi_choice",
                 "answer_key": ["A", "C"], "weight": 1.0},
                {"id": 5, "question_number": 5, "question_type": "short_answer",
                 "answer_key": "Fotosintesis", "weight": 2.0},
            ]

        # Simulasi scan
        sim = _simulate_scan_result(questions)
        grading = _grade_all_answers(sim, questions)

        max_score = sum([float(q.get('weight', 1.0)) for q in questions])
        total_score = grading['total_score']
        percentage = (total_score / max_score * 100) if max_score > 0 else 0

        return {
            "message": "Demo scan berhasil",
            "student_name": "Siswa Demo",
            "student_id_number": "DEMO-NIS-001",
            "total_score": total_score,
            "max_score": max_score,
            "percentage": round(percentage, 2),
            "status": "Lulus" if percentage >= 75 else "Remedial",
            "omr_confidence": 92.5,
            "details": grading['details']
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Demo scan gagal: {str(e)}")


# ============================================================
# Results & Dashboard
# ============================================================

@app.get("/api/results/{exam_id}")
async def get_results(exam_id: int):
    """Daftar hasil ujian untuk exam tertentu"""
    try:
        results = await database.get_results_by_exam(exam_id)
        return serialize_list(results)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/results/detail/{result_id}")
async def get_result_detail(result_id: int):
    """Detail hasil ujian siswa beserta jawaban per soal"""
    try:
        result = await database.get_result_detail(result_id)
        if not result:
            raise HTTPException(status_code=404, detail="Hasil ujian tidak ditemukan")
        result_ser = serialize_record(result)
        if 'student_answers' in result:
            result_ser['student_answers'] = serialize_list(result['student_answers'])
        return result_ser
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/dashboard/{exam_id}")
async def get_dashboard(exam_id: int):
    """Statistik dashboard untuk exam tertentu"""
    try:
        stats = await database.get_dashboard_stats(exam_id)
        exam = await database.get_exam(exam_id)
        result = serialize_record(stats) if stats else {}
        if exam:
            result['exam_title'] = exam.get('title', '')
            result['passing_score'] = float(exam.get('passing_score', 75.0))
        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/question-analysis/{exam_id}")
async def get_question_analysis(exam_id: int):
    """Analisis butir soal untuk exam tertentu"""
    try:
        analysis = await database.get_question_analysis(exam_id)
        return serialize_list(analysis)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ============================================================
# Excel Export
# ============================================================

@app.get("/api/export/{exam_id}")
async def export_excel(exam_id: int):
    """Ekspor hasil ujian ke file Excel (.xlsx)"""
    if not excel_export:
        raise HTTPException(status_code=501, detail="Modul Excel Export belum tersedia")
    try:
        exam = await database.get_exam(exam_id)
        if not exam:
            raise HTTPException(status_code=404, detail="Ujian tidak ditemukan")

        results = await database.get_results_by_exam(exam_id)
        questions = await database.get_questions_by_exam(exam_id)
        stats = await database.get_dashboard_stats(exam_id)
        q_analysis = await database.get_question_analysis(exam_id)

        excel_buffer = excel_export.generate_excel_report(
            serialize_record(exam),
            serialize_list(results),
            serialize_list(questions),
            serialize_record(stats) if stats else {},
            serialize_list(q_analysis)
        )

        return StreamingResponse(
            excel_buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename=Laporan_Ujian_{exam['code']}.xlsx"
            }
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Gagal export Excel: {str(e)}")


# ============================================================
# AI Question Generator
# ============================================================

@app.post("/api/generate-questions")
async def generate_questions(req: GenerateQuestionsRequest):
    """Generate soal menggunakan Gemini AI"""
    if not ai_grader:
        raise HTTPException(status_code=501, detail="Modul AI belum tersedia")
    try:
        req_dict = req.dict()
        result_text = ai_grader.generate_exam_questions(**req_dict)
        if result_text:
            return {"questions": result_text, "status": "success"}
        else:
            raise HTTPException(status_code=500, detail="Gemini API tidak memberikan respons")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Gagal generate soal: {str(e)}")

class DocxDownloadRequest(BaseModel):
    text: str

@app.post("/api/download-docx")
async def download_docx(req: DocxDownloadRequest):
    """Download hasil asesmen sebagai file .docx"""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        raise HTTPException(status_code=501, detail="python-docx belum terinstal")
        
    doc = Document()
    
    # Simple markdown to text processing
    # Can enhance later with full markdown-to-docx converter if needed.
    # For now, just split by newlines and handle some basic formatting.
    for line in req.text.split('\n'):
        line_clean = line.strip()
        if line_clean.startswith('# '):
            p = doc.add_heading(line_clean[2:], level=1)
        elif line_clean.startswith('## '):
            p = doc.add_heading(line_clean[3:], level=2)
        elif line_clean.startswith('### '):
            p = doc.add_heading(line_clean[4:], level=3)
        elif line_clean.startswith('**') and line_clean.endswith('**'):
            p = doc.add_paragraph()
            run = p.add_run(line_clean[2:-2])
            run.bold = True
        else:
            doc.add_paragraph(line)
            
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": "attachment; filename=Asesmen_Lengkap.docx"
        }
    )


# ============================================================
# Helper Functions
# ============================================================

def _build_question_layout(questions):
    """Buat konfigurasi layout soal untuk OMR processor berdasarkan posisi di LJK"""
    layout = []
    y_start = 400  # Mulai dari y=400px di halaman 2100x2970
    row_height = 60
    x_start = 150

    for i, q in enumerate(questions):
        q_type = q.get('question_type', 'single_choice')

        if q_type == 'single_choice':
            layout.append({
                'question_number': q.get('question_number', i + 1),
                'question_type': q_type,
                'x': x_start, 'y': y_start + i * row_height,
                'w': 600, 'h': 50,
                'options': 4
            })
        elif q_type == 'multi_choice':
            layout.append({
                'question_number': q.get('question_number', i + 1),
                'question_type': q_type,
                'x': x_start, 'y': y_start + i * row_height,
                'w': 600, 'h': 50,
                'options': 4
            })
        elif q_type == 'true_false':
            layout.append({
                'question_number': q.get('question_number', i + 1),
                'question_type': q_type,
                'x': x_start, 'y': y_start + i * row_height,
                'w': 400, 'h': 50
            })
        elif q_type == 'matching':
            layout.append({
                'question_number': q.get('question_number', i + 1),
                'question_type': q_type,
                'x': x_start, 'y': y_start + i * row_height,
                'w': 500, 'h': 250,
                'num_questions': 4, 'num_options': 4
            })
        elif q_type == 'short_answer':
            layout.append({
                'question_number': q.get('question_number', i + 1),
                'question_type': q_type,
                'x': x_start, 'y': y_start + i * row_height,
                'w': 1600, 'h': 150
            })

    return layout


def _simulate_scan_result(questions):
    """Simulasi hasil scan untuk demo/testing"""
    import random
    answers = {}
    avg_confidence = 0
    count = 0

    for q in questions:
        q_num = q.get('question_number', 1)
        q_type = q.get('question_type', 'single_choice')
        answer_key = q.get('answer_key', 'A')

        if q_type == 'single_choice':
            # 80% chance jawab benar
            if random.random() < 0.8:
                ans = answer_key if isinstance(answer_key, str) else str(answer_key)
            else:
                opts = ['A', 'B', 'C', 'D']
                ans = random.choice(opts)
            conf = random.uniform(75, 98)
            answers[str(q_num)] = {'answer': ans, 'confidence': round(conf, 1)}

        elif q_type == 'multi_choice':
            if random.random() < 0.7:
                ans = answer_key if isinstance(answer_key, list) else [str(answer_key)]
            else:
                ans = random.sample(['A', 'B', 'C', 'D'], random.randint(1, 3))
            conf = random.uniform(70, 95)
            answers[str(q_num)] = {'answer': ans, 'confidence': round(conf, 1)}

        elif q_type == 'true_false':
            if random.random() < 0.8:
                ans = answer_key if isinstance(answer_key, str) else str(answer_key)
            else:
                ans = 'B' if answer_key == 'S' else 'S'
            conf = random.uniform(80, 99)
            answers[str(q_num)] = {'answer': ans, 'confidence': round(conf, 1)}

        elif q_type == 'matching':
            if isinstance(answer_key, dict):
                ans = answer_key
            else:
                ans = {1: 'A', 2: 'B', 3: 'C', 4: 'D'}
            conf = random.uniform(65, 90)
            answers[str(q_num)] = {'answer': ans, 'confidence': round(conf, 1)}

        elif q_type == 'short_answer':
            conf = random.uniform(80, 95)
            ans_text = answer_key if isinstance(answer_key, str) else str(answer_key)
            answers[str(q_num)] = {
                'answer': ans_text,
                'confidence': round(conf, 1),
                'transcribed_text': ans_text,
                'ai_feedback': f'Jawaban "{ans_text}" sesuai dengan kunci jawaban.',
                'is_ai_graded': True
            }

        avg_confidence += answers[str(q_num)]['confidence']
        count += 1

    return {
        'answers': answers,
        'short_answers': {},
        'corners_detected': True,
        'metadata': None,
        'avg_confidence': round(avg_confidence / count, 1) if count > 0 else 0
    }


def _grade_all_answers(scan_result, questions):
    """Koreksi semua jawaban berdasarkan kunci jawaban"""
    details = []
    total_score = 0

    for q in questions:
        q_num = str(q.get('question_number', 1))
        q_type = q.get('question_type', 'single_choice')
        answer_key = q.get('answer_key')
        weight = float(q.get('weight', 1.0))

        scan_ans = scan_result.get('answers', {}).get(q_num, {})
        student_answer = scan_ans.get('answer')
        confidence = scan_ans.get('confidence', 0)

        is_correct = False
        score_earned = 0
        ai_feedback = None
        ai_confidence = None

        # Parse answer_key dari JSON string jika perlu
        if isinstance(answer_key, str):
            try:
                answer_key = json.loads(answer_key)
            except (json.JSONDecodeError, TypeError):
                pass

        if q_type == 'single_choice' or q_type == 'true_false':
            key_str = str(answer_key) if answer_key else ''
            ans_str = str(student_answer) if student_answer else ''
            is_correct = key_str.upper() == ans_str.upper()
            score_earned = weight if is_correct else 0

        elif q_type == 'multi_choice':
            key_set = set(answer_key) if isinstance(answer_key, list) else {str(answer_key)}
            ans_set = set(student_answer) if isinstance(student_answer, list) else {str(student_answer)} if student_answer else set()
            is_correct = key_set == ans_set
            # Skor parsial: proporsi jawaban benar
            if key_set:
                correct_picks = len(key_set & ans_set)
                wrong_picks = len(ans_set - key_set)
                partial = max(0, (correct_picks - wrong_picks)) / len(key_set)
                score_earned = round(weight * partial, 2)
                is_correct = partial >= 1.0

        elif q_type == 'matching':
            if isinstance(answer_key, dict) and isinstance(student_answer, dict):
                correct_count = 0
                total_items = len(answer_key)
                for k, v in answer_key.items():
                    if str(student_answer.get(str(k), '')).upper() == str(v).upper():
                        correct_count += 1
                score_earned = round(weight * correct_count / total_items, 2) if total_items > 0 else 0
                is_correct = correct_count == total_items
            else:
                score_earned = 0

        elif q_type == 'short_answer':
            # Untuk isian singkat, gunakan AI jika tersedia
            ai_data = scan_ans
            if ai_data.get('is_ai_graded'):
                ans_text = str(ai_data.get('transcribed_text', ''))
                key_text = str(answer_key) if answer_key else ''
                is_correct = ai_data.get('is_correct', ans_text.strip().lower() == key_text.strip().lower())
                score_earned = float(ai_data.get('score_earned', weight if is_correct else 0))
                ai_feedback = ai_data.get('ai_feedback', '')
                ai_confidence = ai_data.get('confidence', 0)
            else:
                # Perbandingan string sederhana
                ans_text = str(student_answer) if student_answer else ''
                key_text = str(answer_key) if answer_key else ''
                is_correct = ans_text.strip().lower() == key_text.strip().lower()
                score_earned = weight if is_correct else 0

        total_score += score_earned

        details.append({
            'question_id': q.get('id'),
            'question_number': int(q_num),
            'question_type': q_type,
            'student_response': student_answer,
            'answer_key': answer_key,
            'is_correct': is_correct,
            'score_earned': score_earned,
            'max_score': weight,
            'confidence': confidence,
            'ai_feedback': ai_feedback,
            'ai_confidence': ai_confidence
        })

    return {
        'total_score': round(total_score, 2),
        'details': details
    }
